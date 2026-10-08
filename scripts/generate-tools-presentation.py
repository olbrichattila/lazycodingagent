#!/usr/bin/env python3
"""Generate LazyCodingAgent agent tools overview as PowerPoint (.pptx) and PDF."""

from __future__ import annotations

import sys
from pathlib import Path

try:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches, Pt
except ImportError:
    print("Missing dependency: pip install python-pptx", file=sys.stderr)
    sys.exit(1)

REPO_ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = REPO_ROOT / "presentations"
PPTX_PATH = OUTPUT_DIR / "LazyCodingAgent-Agent-Tools.pptx"
PDF_PATH = OUTPUT_DIR / "LazyCodingAgent-Agent-Tools.pdf"

# Metadata aligned with uTool*.pas GetDefinition and Execute output shapes.
TOOLS = [
    {
        "name": "read_file",
        "aliases": [],
        "modes": "Ask, Plan, Agent",
        "description": (
            "Lets the agent inspect project files without executing code. "
            "Returns line-numbered text so it can cite locations; use offset and "
            "limit to read large files in chunks (up to 2000 lines, 16 MiB cap)."
        ),
        "inputs": [
            "path (string) — required",
            "offset (integer) — optional, 1-based start line",
            "limit (integer) — optional, max lines (default 2000)",
        ],
        "outputs": [
            "path, content (line-numbered), offset, lines_returned, total_lines",
            "On failure: error message in JSON",
        ],
    },
    {
        "name": "write_file",
        "aliases": [],
        "modes": "Agent only",
        "description": (
            "Creates a new file or replaces an entire file in one step. "
            "The agent sends the full new content; missing parent folders are "
            "created automatically under the workspace."
        ),
        "inputs": [
            "path (string) — required",
            "content (string) — required",
        ],
        "outputs": [
            "status, path, bytes_written, changed_paths",
            "Optional: tracking_warning",
            "On failure: error",
        ],
    },
    {
        "name": "list_directory",
        "aliases": [],
        "modes": "Ask, Plan, Agent",
        "description": (
            "Explores the project one directory at a time, like a single folder "
            "in a file browser. Lists immediate children (files, folders, symlinks) "
            "sorted by name, up to 500 entries per call."
        ),
        "inputs": [
            "path (string) — optional, default project root",
            "include_hidden (boolean) — optional",
        ],
        "outputs": [
            "directory, entries[] (path, type, size), total_entries, truncated",
            "On failure: error",
        ],
    },
    {
        "name": "glob",
        "aliases": [],
        "modes": "Ask, Plan, Agent",
        "description": (
            "Finds many matching paths under a folder using wildcards (*, ?, **). "
            "Use when the agent knows a naming pattern (e.g. all units under src/) "
            "instead of walking directories one level at a time."
        ),
        "inputs": [
            "pattern (string) — required, e.g. src/**/*",
            "path (string) — optional search root",
            "include_hidden (boolean) — optional",
        ],
        "outputs": [
            "directory, entries[] (path, type, size), total_entries, truncated",
            "On failure: error",
        ],
    },
    {
        "name": "search_code",
        "aliases": ["grep"],
        "modes": "Ask, Plan, Agent",
        "description": (
            "Searches file contents across the workspace, similar to IDE Find in Files. "
            "Supports literal text or regex, optional path and glob filters; skips "
            "oversized files, binary data, and returns up to 500 matches."
        ),
        "inputs": [
            "query (string) — required",
            "path (string) — optional file or directory",
            "glob (string) — optional path filter",
            "regex, case_sensitive, include_hidden (boolean) — optional",
        ],
        "outputs": [
            "matches[] (path, line, column, text), total_matches, truncated",
            "On failure: error",
        ],
    },
    {
        "name": "apply_patch",
        "aliases": [],
        "modes": "Agent only",
        "description": (
            "Applies multi-line edits using standard unified diff format. "
            "The agent validates every hunk before writing; can create, modify, "
            "or delete files—preferred for structured code changes."
        ),
        "inputs": [
            "patch (string) — required unified diff text",
        ],
        "outputs": [
            "status, changed_paths",
            "Optional: warnings",
            "On failure: error, error_kind, hint; may include partial",
        ],
    },
    {
        "name": "edit_file",
        "aliases": [],
        "modes": "Agent only",
        "description": (
            "Performs one exact search-and-replace in a single file. "
            "Best for a small, localized change when a full patch is awkward; "
            "old_text must match byte-for-byte exactly once in the file."
        ),
        "inputs": [
            "path (string) — required",
            "old_text (string) — required, exact match once",
            "new_text (string) — required",
        ],
        "outputs": [
            "status, changed_paths",
            "Optional: tracking_warning, warnings",
            "On failure: error",
        ],
    },
    {
        "name": "shell",
        "aliases": ["terminal", "run_process"],
        "modes": "Agent only",
        "description": (
            "Runs terminal commands such as builds, tests, and package installs. "
            "Executes under the project root (or optional cwd) with a timeout; "
            "captures stdout/stderr and may refresh the workspace file tree after run."
        ),
        "inputs": [
            "command (string) — required",
            "cwd (string) — optional working directory",
            "timeout_ms (integer) — optional (default 120000)",
        ],
        "outputs": [
            "exit_code, stdout, stderr, timed_out, cancelled",
            "stdout_truncated, stderr_truncated, changed_paths, refresh_project",
            "On failure: error",
        ],
    },
    {
        "name": "diagnostics",
        "aliases": [],
        "modes": "Ask, Plan, Agent (build action: Agent only)",
        "description": (
            "Shows compiler and linter messages so the agent can fix build errors. "
            "Read returns the last cached result; in Agent mode, build runs an "
            "auto-detected or custom command and parses output into structured messages."
        ),
        "inputs": [
            "action (string) — required: read or build",
            "target (string) — optional project/manifest/source",
            "timeout_ms (integer) — optional for build",
            "command (string) — optional custom build command",
        ],
        "outputs": [
            "read: cached diagnostics JSON or status unavailable",
            "build: status, messages[], runner, target, exit_code, stdout, stderr",
            "On failure: error",
        ],
    },
    {
        "name": "git",
        "aliases": [],
        "modes": "Ask, Plan, Agent",
        "description": (
            "Read-only view of Git history and working tree changes. "
            "The agent can check status, diffs, log, and show commits—without "
            "committing, pushing, or other mutating Git operations."
        ),
        "inputs": [
            "operation (string) — required: status | diff | log | show",
            "revision (string) — optional",
            "paths (array of strings) — optional",
            "limit (integer) — optional, 1–100 for log",
            "staged (boolean) — optional for diff",
        ],
        "outputs": [
            "exit_code, stdout, stderr",
            "Optional: error if Git failed",
            "On failure: error",
        ],
    },
    {
        "name": "todo",
        "aliases": ["plan"],
        "modes": "Ask, Plan, Agent",
        "description": (
            "Maintains a task checklist for the current chat session only. "
            "Helps the agent track what is pending, in progress, or done while "
            "working on a multi-step request (not saved to project files)."
        ),
        "inputs": [
            "action (string) — required: read or replace",
            "items (array) — required for replace: id, text, status per task",
        ],
        "outputs": [
            "items[] (id, text, status)",
            "On failure: error",
        ],
    },
    {
        "name": "create_plan_file",
        "aliases": [],
        "modes": "Plan, Agent",
        "description": (
            "Saves a finished implementation plan as Markdown in the workspace .plan/ "
            "folder. Used in Plan mode so the user can review steps before Agent mode "
            "applies code changes; each file gets a unique timestamped name."
        ),
        "inputs": [
            "content (string) — required, markdown plan (max 1 MiB)",
        ],
        "outputs": [
            "status, path, content",
            "On failure: error",
        ],
    },
]

# Shared theme (PDF + PowerPoint).
RGB_VIOLET = RGBColor(0x5B, 0x21, 0xB6)
RGB_PINK = RGBColor(0xDB, 0x27, 0x77)
RGB_CARD_FILL = RGBColor(0xFD, 0xF4, 0xFF)
RGB_CARD_STROKE = RGBColor(0xE9, 0xD5, 0xFF)
RGB_TEXT_MID = RGBColor(0x58, 0x1C, 0x87)
RGB_ACCENT = RGBColor(0x93, 0x33, 0xEA)
RGB_CHIP_BG = RGBColor(0xF3, 0xE8, 0xFF)
RGB_WHITE = RGBColor(0xFF, 0xFF, 0xFF)
RGB_ICON_RING = RGBColor(0x7E, 0x22, 0xCE)
RGB_LILAC = RGBColor(0xC0, 0x84, 0xFC)
RGB_BULLET = RGBColor(0xA8, 0x55, 0xF7)
RGB_SUBTITLE = RGBColor(0xFA, 0xE8, 0xFF)

TOOL_ICON_KIND: dict[str, str] = {
    "read_file": "doc",
    "write_file": "pen",
    "list_directory": "folder",
    "glob": "globe",
    "search_code": "search",
    "apply_patch": "patch",
    "edit_file": "edit",
    "shell": "terminal",
    "diagnostics": "pulse",
    "git": "branch",
    "todo": "checklist",
    "create_plan_file": "plan",
}

PPTX_CARD_L = Inches(0.5)
PPTX_CARD_T = Inches(0.5)
PPTX_CARD_W = Inches(9.0)
PPTX_CARD_H = Inches(6.5)
PPTX_PAD = Inches(0.4)
PPTX_HEADER_H = Inches(0.95)


def _icon_kind_for_slide_title(title: str) -> str:
    if title in TOOL_ICON_KIND:
        return TOOL_ICON_KIND[title]
    for tool in TOOLS:
        if title in tool["aliases"]:
            return TOOL_ICON_KIND[tool["name"]]
    return "agent"


def _pptx_solid_fill(shape, rgb: RGBColor, line_rgb: RGBColor | None = None) -> None:
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb
    if line_rgb is not None:
        shape.line.color.rgb = line_rgb
        shape.line.width = Pt(1.2)
    else:
        shape.line.fill.background()


def _pptx_slide_gradient(slide) -> None:
    fill = slide.background.fill
    fill.gradient()
    fill.gradient_angle = 270
    stops = fill.gradient_stops
    stops[0].color.rgb = RGB_VIOLET
    stops[1].color.rgb = RGB_PINK


def _pptx_style_run(run, size_pt: int, *, bold=False, italic=False, color=RGB_TEXT_MID):
    run.font.name = "Calibri"
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color


def _pptx_add_chip(slide, left, top, label: str, width=None) -> None:
    chip_h = Inches(0.28)
    if width is None:
        width = Inches(0.15 * len(label) + 0.55)
    chip = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, chip_h
    )
    _pptx_solid_fill(chip, RGB_CHIP_BG, RGB_CARD_STROKE)
    chip.adjustments[0] = 0.15
    tf = chip.text_frame
    tf.margin_left = Pt(8)
    tf.margin_top = Pt(2)
    tf.word_wrap = False
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = label
    _pptx_style_run(run, 13, bold=True, color=RGB_ACCENT)


def _pptx_add_bullets(slide, left, top, width, height, items: list[str]) -> None:
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = Pt(2)
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(8)
        dot = p.add_run()
        dot.text = "● "
        _pptx_style_run(dot, 13, color=RGB_BULLET)
        body = p.add_run()
        body.text = item
        _pptx_style_run(body, 13, color=RGB_TEXT_MID)


def _pptx_draw_icon(slide, kind: str, cx, cy, diameter) -> None:
    r = diameter / 2
    ring = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, cx - r, cy - r, diameter, diameter
    )
    _pptx_solid_fill(ring, RGB_ICON_RING)
    s = diameter * 0.55
    ix, iy = cx - s / 2, cy - s / 2
    white = RGB_WHITE
    if kind in ("doc", "plan", "terminal", "folder", "patch"):
        rect = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, ix, iy, s, s * 0.85)
        _pptx_solid_fill(rect, white)
        rect.adjustments[0] = 0.08
    elif kind == "pen":
        rect = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, cx - s * 0.12, iy, s * 0.24, s * 0.75
        )
        _pptx_solid_fill(rect, white)
    elif kind in ("globe", "agent", "search"):
        circ = slide.shapes.add_shape(MSO_SHAPE.OVAL, ix, iy, s, s)
        circ.fill.background()
        circ.line.color.rgb = white
        circ.line.width = Pt(2)
    elif kind == "edit":
        line = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, ix, iy + s * 0.4, s, Pt(3)
        )
        _pptx_solid_fill(line, white)
        line.rotation = 315
    elif kind in ("pulse", "branch", "checklist"):
        circ = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, cx - s * 0.18, cy - s * 0.18, s * 0.36, s * 0.36
        )
        _pptx_solid_fill(circ, white)
    else:
        circ = slide.shapes.add_shape(MSO_SHAPE.OVAL, ix, iy, s, s)
        _pptx_solid_fill(circ, white)


def _add_styled_pptx_slide(prs: Presentation, data: dict) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _pptx_slide_gradient(slide)

    card = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, PPTX_CARD_L, PPTX_CARD_T, PPTX_CARD_W, PPTX_CARD_H
    )
    _pptx_solid_fill(card, RGB_CARD_FILL, RGB_CARD_STROKE)
    card.adjustments[0] = 0.04

    header_top = PPTX_CARD_T
    header = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        PPTX_CARD_L,
        header_top,
        PPTX_CARD_W,
        PPTX_HEADER_H,
    )
    _pptx_solid_fill(header, RGB_ACCENT)
    header_top_cap = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        PPTX_CARD_L,
        header_top,
        PPTX_CARD_W,
        PPTX_HEADER_H,
    )
    _pptx_solid_fill(header_top_cap, RGB_ACCENT)
    header_top_cap.adjustments[0] = 0.12

    icon_d = Inches(0.58)
    icon_cx = PPTX_CARD_L + PPTX_PAD + icon_d / 2
    icon_cy = header_top + PPTX_HEADER_H / 2
    _pptx_draw_icon(slide, _icon_kind_for_slide_title(data["title"]), icon_cx, icon_cy, icon_d)

    title_l = PPTX_CARD_L + PPTX_PAD + Inches(0.72)
    title_t = header_top + Inches(0.22)
    title_w = PPTX_CARD_W - PPTX_PAD * 2 - Inches(0.72)
    title_size = 30 if data.get("is_title") else 28
    title_box = slide.shapes.add_textbox(title_l, title_t, title_w, Inches(0.55))
    tr = title_box.text_frame.paragraphs[0].add_run()
    tr.text = data["title"]
    _pptx_style_run(tr, title_size, bold=True, color=RGB_WHITE)

    body_top = header_top + PPTX_HEADER_H + PPTX_PAD
    content_l = PPTX_CARD_L + PPTX_PAD
    content_w = PPTX_CARD_W - PPTX_PAD * 2

    if data.get("subtitle"):
        sub_box = slide.shapes.add_textbox(
            title_l, title_t + Inches(0.48), title_w, Inches(0.35)
        )
        sr = sub_box.text_frame.paragraphs[0].add_run()
        sr.text = data["subtitle"]
        _pptx_style_run(sr, 14, italic=True, color=RGB_SUBTITLE)

    if data.get("is_title"):
        _pptx_add_chip(slide, content_l, body_top, "Overview")
        desc_box = slide.shapes.add_textbox(
            content_l, body_top + Inches(0.38), content_w, Inches(4.5)
        )
        dr = desc_box.text_frame.paragraphs[0].add_run()
        dr.text = data["description"]
        _pptx_style_run(dr, 18, color=RGB_TEXT_MID)
        return

    y = body_top
    _pptx_add_chip(slide, content_l, y, "What it does")
    y += Inches(0.36)
    desc_box = slide.shapes.add_textbox(content_l, y, content_w, Inches(1.35))
    dtf = desc_box.text_frame
    dtf.word_wrap = True
    dr = dtf.paragraphs[0].add_run()
    dr.text = data["description"]
    _pptx_style_run(dr, 14, color=RGB_TEXT_MID)

    y += Inches(1.42)
    col_gap = Inches(0.2)
    col_w = (content_w - col_gap) / 2
    col2_l = content_l + col_w + col_gap
    _pptx_add_chip(slide, content_l, y, "Input", width=col_w * 0.45)
    _pptx_add_chip(slide, col2_l, y, "Output", width=col_w * 0.5)
    list_top = y + Inches(0.34)
    list_h = Inches(2.35)
    _pptx_add_bullets(slide, content_l, list_top, col_w, list_h, data["inputs"])
    _pptx_add_bullets(slide, col2_l, list_top, col_w, list_h, data["outputs"])

    modes_top = PPTX_CARD_T + PPTX_CARD_H - PPTX_PAD - Inches(0.32)
    modes_label = f"Modes: {data['modes']}"
    _pptx_add_chip(
        slide,
        content_l,
        modes_top,
        modes_label,
        width=min(content_w, Inches(0.12 * len(modes_label) + 0.35)),
    )


def _iter_slides():
    yield {
        "title": "LazyCodingAgent — Agent Tools",
        "subtitle": None,
        "description": (
            "These are the functions the LazyCodingAgent LLM can call while helping "
            "you in the workspace. Each call uses JSON arguments; the app returns "
            "JSON results (errors include an error field)."
        ),
        "inputs": [],
        "outputs": [],
        "modes": "",
        "is_title": True,
    }
    for tool in TOOLS:
        yield {
            "title": tool["name"],
            "subtitle": None,
            "description": tool["description"],
            "inputs": tool["inputs"],
            "outputs": tool["outputs"],
            "modes": tool["modes"],
            "is_title": False,
        }
    for tool in TOOLS:
        primary = tool["name"]
        for alias in tool["aliases"]:
            yield {
                "title": alias,
                "subtitle": (
                    f"Alias for {primary} — same behavior, inputs, and outputs."
                ),
                "description": tool["description"],
                "inputs": tool["inputs"],
                "outputs": tool["outputs"],
                "modes": tool["modes"],
                "is_title": False,
            }


def build_presentation() -> Presentation:
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)
    for slide_data in _iter_slides():
        _add_styled_pptx_slide(prs, slide_data)
    return prs


def build_pdf(path: Path) -> int:
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import landscape, A4
        from reportlab.lib.units import inch
        from reportlab.pdfgen import canvas
    except ImportError:
        print("Missing dependency: pip install reportlab", file=sys.stderr)
        return 1

    page_w, page_h = landscape(A4)
    margin = 0.5 * inch
    card_pad = 0.4 * inch
    line_h_body = 20
    line_h_tight = 18
    fs_title = 28
    fs_title_cover = 30
    fs_subtitle = 14
    fs_section = 13
    fs_body = 14
    fs_bullet = 13
    fs_cover_body = 18

    violet = colors.HexColor("#5b21b6")
    pink = colors.HexColor("#db2777")
    lilac = colors.HexColor("#c084fc")
    card_fill = colors.HexColor("#fdf4ff")
    card_stroke = colors.HexColor("#e9d5ff")
    text_dark = colors.HexColor("#3b0764")
    text_mid = colors.HexColor("#581c87")
    accent = colors.HexColor("#9333ea")
    chip_bg = colors.HexColor("#f3e8ff")
    bullet = colors.HexColor("#a855f7")

    def draw_gradient_background(c: canvas.Canvas) -> None:
        steps = 48
        for i in range(steps):
            t = i / max(steps - 1, 1)
            r = violet.red + (pink.red - violet.red) * t
            g = violet.green + (pink.green - violet.green) * t
            b = violet.blue + (pink.blue - violet.blue) * t
            c.setFillColorRGB(r, g, b)
            band_h = page_h / steps + 1
            c.rect(0, page_h - (i + 1) * band_h, page_w, band_h, fill=1, stroke=0)

    def draw_icon(
        c: canvas.Canvas, kind: str, cx: float, cy: float, size: float
    ) -> None:
        s = size
        c.setStrokeColor(colors.white)
        c.setFillColor(colors.white)
        c.setLineWidth(2)
        if kind == "agent":
            c.circle(cx, cy, s * 0.38, fill=1, stroke=1)
            c.setFillColor(lilac)
            c.circle(cx - s * 0.12, cy + s * 0.06, s * 0.06, fill=1, stroke=0)
            c.circle(cx + s * 0.12, cy + s * 0.06, s * 0.06, fill=1, stroke=0)
            c.setStrokeColor(lilac)
            c.setLineWidth(2.5)
            c.arc(cx - s * 0.18, cy - s * 0.22, cx + s * 0.18, cy - s * 0.02, 200, 340)
        elif kind == "doc":
            c.roundRect(cx - s * 0.28, cy - s * 0.34, s * 0.56, s * 0.68, 4, fill=1, stroke=1)
            c.setStrokeColor(lilac)
            for i in range(3):
                y = cy + s * 0.12 - i * s * 0.14
                c.line(cx - s * 0.14, y, cx + s * 0.14, y)
        elif kind == "pen":
            c.setFillColor(colors.white)
            c.rect(cx - s * 0.08, cy - s * 0.32, s * 0.16, s * 0.52, fill=1, stroke=1)
            c.setFillColor(lilac)
            path = c.beginPath()
            path.moveTo(cx - s * 0.08, cy + s * 0.2)
            path.lineTo(cx, cy + s * 0.36)
            path.lineTo(cx + s * 0.08, cy + s * 0.2)
            path.close()
            c.drawPath(path, fill=1, stroke=0)
        elif kind == "folder":
            c.roundRect(cx - s * 0.32, cy - s * 0.28, s * 0.64, s * 0.48, 5, fill=1, stroke=1)
            c.roundRect(cx - s * 0.28, cy + s * 0.02, s * 0.32, s * 0.12, 3, fill=1, stroke=1)
        elif kind == "globe":
            c.circle(cx, cy, s * 0.3, fill=1, stroke=1)
            c.setStrokeColor(lilac)
            c.line(cx - s * 0.3, cy, cx + s * 0.3, cy)
            c.circle(cx, cy, s * 0.18, fill=0, stroke=1)
        elif kind == "search":
            c.circle(cx - s * 0.08, cy + s * 0.06, s * 0.18, fill=0, stroke=1)
            c.line(cx + s * 0.06, cy - s * 0.1, cx + s * 0.28, cy - s * 0.32)
        elif kind == "patch":
            c.setFillColor(colors.white)
            c.rect(cx - s * 0.26, cy - s * 0.08, s * 0.52, s * 0.16, fill=1, stroke=1)
            c.setFillColor(lilac)
            c.rect(cx - s * 0.1, cy - s * 0.22, s * 0.2, s * 0.14, fill=1, stroke=0)
        elif kind == "edit":
            c.line(cx - s * 0.22, cy + s * 0.22, cx + s * 0.22, cy - s * 0.22)
            c.line(cx + s * 0.08, cy - s * 0.22, cx + s * 0.22, cy - s * 0.06)
            c.line(cx + s * 0.22, cy - s * 0.22, cx + s * 0.08, cy - s * 0.22)
        elif kind == "terminal":
            c.roundRect(cx - s * 0.32, cy - s * 0.26, s * 0.64, s * 0.52, 5, fill=1, stroke=1)
            c.setStrokeColor(lilac)
            c.line(cx - s * 0.18, cy + s * 0.04, cx - s * 0.06, cy - s * 0.04)
            c.line(cx - s * 0.06, cy - s * 0.04, cx - s * 0.18, cy - s * 0.12)
            c.line(cx - s * 0.02, cy - s * 0.12, cx + s * 0.18, cy - s * 0.12)
        elif kind == "pulse":
            c.setStrokeColor(colors.white)
            c.line(cx - s * 0.28, cy, cx - s * 0.12, cy)
            c.line(cx - s * 0.12, cy, cx - s * 0.04, cy + s * 0.18)
            c.line(cx - s * 0.04, cy + s * 0.18, cx + s * 0.08, cy - s * 0.2)
            c.line(cx + s * 0.08, cy - s * 0.2, cx + s * 0.28, cy)
        elif kind == "branch":
            c.circle(cx - s * 0.12, cy + s * 0.16, s * 0.07, fill=1, stroke=1)
            c.circle(cx + s * 0.12, cy + s * 0.16, s * 0.07, fill=1, stroke=1)
            c.circle(cx, cy - s * 0.14, s * 0.07, fill=1, stroke=1)
            c.line(cx - s * 0.12, cy + s * 0.09, cx, cy - s * 0.07)
            c.line(cx + s * 0.12, cy + s * 0.09, cx, cy - s * 0.07)
        elif kind == "checklist":
            for i in range(3):
                y = cy + s * 0.16 - i * s * 0.16
                c.roundRect(cx - s * 0.22, y - s * 0.06, s * 0.1, s * 0.1, 2, fill=0, stroke=1)
                c.line(cx - s * 0.06, y, cx + s * 0.22, y)
        elif kind == "plan":
            c.roundRect(cx - s * 0.26, cy - s * 0.3, s * 0.52, s * 0.6, 4, fill=1, stroke=1)
            c.setStrokeColor(lilac)
            c.line(cx - s * 0.14, cy + s * 0.14, cx + s * 0.14, cy + s * 0.14)
            c.line(cx - s * 0.14, cy, cx + s * 0.06, cy)
            c.line(cx - s * 0.14, cy - s * 0.14, cx + s * 0.1, cy - s * 0.14)
        else:
            c.circle(cx, cy, s * 0.28, fill=1, stroke=1)

    def draw_wrapped(
        c: canvas.Canvas,
        text: str,
        x: float,
        y: float,
        max_w: float,
        font: str,
        size: int,
        color=text_dark,
        leading: int = line_h_body,
    ) -> float:
        c.setFillColor(color)
        c.setFont(font, size)
        words = text.split()
        line = ""
        for word in words:
            trial = (line + " " + word).strip()
            if c.stringWidth(trial, font, size) <= max_w:
                line = trial
            else:
                if line:
                    c.drawString(x, y, line)
                    y -= leading
                line = word
        if line:
            c.drawString(x, y, line)
            y -= leading
        return y

    def draw_section_label(
        c: canvas.Canvas, label: str, x: float, y: float, width: float | None = None
    ) -> float:
        c.setFillColor(chip_bg)
        c.setStrokeColor(card_stroke)
        w = (width if width is not None else c.stringWidth(label, "Helvetica-Bold", fs_section) + 20)
        c.roundRect(x, y - 4, w, 22, 8, fill=1, stroke=1)
        c.setFillColor(accent)
        c.setFont("Helvetica-Bold", fs_section)
        c.drawString(x + 10, y + 2, label)
        return y - 28

    def draw_bullets(
        c: canvas.Canvas,
        items: list[str],
        x: float,
        y: float,
        max_w: float,
        start_y: float | None = None,
    ) -> float:
        if start_y is not None:
            y = start_y
        for item in items:
            c.setFillColor(bullet)
            c.circle(x + 5, y + 5, 3.5, fill=1, stroke=0)
            y = draw_wrapped(
                c,
                item,
                x + 16,
                y,
                max_w - 16,
                "Helvetica",
                fs_bullet,
                text_mid,
                line_h_tight,
            )
            y -= 5
        return y

    c = canvas.Canvas(str(path), pagesize=(page_w, page_h))
    card_x = margin
    card_y = margin
    card_w = page_w - 2 * margin
    card_h = page_h - 2 * margin
    content_x = card_x + card_pad
    content_w = card_w - 2 * card_pad

    for slide in _iter_slides():
        draw_gradient_background(c)

        c.setFillColor(card_fill)
        c.setStrokeColor(card_stroke)
        c.setLineWidth(1.2)
        c.roundRect(card_x, card_y, card_w, card_h, 18, fill=1, stroke=1)

        header_h = 0.95 * inch
        c.setFillColor(accent)
        c.roundRect(card_x, card_y + card_h - header_h, card_w, header_h, 18, fill=1, stroke=0)
        c.rect(card_x, card_y + card_h - header_h, card_w, header_h - 18, fill=1, stroke=0)

        icon_kind = _icon_kind_for_slide_title(slide["title"])
        icon_cx = card_x + card_pad + 0.28 * inch
        icon_cy = card_y + card_h - header_h * 0.52
        c.setFillColor(colors.HexColor("#7e22ce"))
        c.circle(icon_cx, icon_cy, 0.32 * inch, fill=1, stroke=0)
        draw_icon(c, icon_kind, icon_cx, icon_cy, 0.42 * inch)

        title_x = card_x + card_pad + 0.72 * inch
        title_y = card_y + card_h - header_h * 0.38
        c.setFillColor(colors.white)
        c.setFont(
            "Helvetica-Bold",
            fs_title_cover if slide.get("is_title") else fs_title,
        )
        c.drawString(title_x, title_y, slide["title"])

        if slide.get("subtitle"):
            c.setFont("Helvetica-Oblique", fs_subtitle)
            c.setFillColor(colors.HexColor("#fae8ff"))
            c.drawString(title_x, title_y - 24, slide["subtitle"])

        y = card_y + card_h - header_h - card_pad - 6
        if slide.get("is_title"):
            y = draw_wrapped(
                c,
                slide["description"],
                content_x,
                y,
                content_w,
                "Helvetica",
                fs_cover_body,
                text_mid,
                24,
            )
        else:
            y = draw_section_label(c, "What it does", content_x, y)
            y = draw_wrapped(
                c,
                slide["description"],
                content_x,
                y,
                content_w,
                "Helvetica",
                fs_body,
                text_mid,
                line_h_body,
            )
            y -= 12

            col_gap = 16
            col_w = (content_w - col_gap) / 2
            col_right = content_x + col_w + col_gap
            labels_y = y
            y = draw_section_label(c, "Input", content_x, labels_y, col_w)
            draw_section_label(c, "Output", col_right, labels_y, col_w)
            bullets_y = y
            y_left = draw_bullets(
                c, slide["inputs"], content_x, bullets_y, col_w, start_y=bullets_y
            )
            y_right = draw_bullets(
                c, slide["outputs"], col_right, bullets_y, col_w, start_y=bullets_y
            )
            y = min(y_left, y_right) - 10

            c.setFillColor(chip_bg)
            c.setStrokeColor(card_stroke)
            modes_text = f"Modes: {slide['modes']}"
            modes_w = min(
                content_w,
                c.stringWidth(modes_text, "Helvetica-Bold", fs_section) + 28,
            )
            c.roundRect(content_x, y - 24, modes_w, 26, 8, fill=1, stroke=1)
            c.setFillColor(accent)
            c.setFont("Helvetica-Bold", fs_section)
            c.drawString(content_x + 12, y - 16, modes_text)

        c.showPage()

    c.save()
    return 0


def main() -> int:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    prs = build_presentation()
    prs.save(str(PPTX_PATH))
    expected = 1 + len(TOOLS) + sum(len(t["aliases"]) for t in TOOLS)
    print(f"Wrote {PPTX_PATH}")
    print(f"Slides: {len(prs.slides)} (expected {expected})")
    if len(prs.slides) != expected:
        print("Warning: slide count mismatch", file=sys.stderr)
        return 1

    if build_pdf(PDF_PATH) != 0:
        return 1
    print(f"Wrote {PDF_PATH}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
